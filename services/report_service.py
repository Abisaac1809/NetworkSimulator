import datetime


class ReportService:
    def __init__(self, config, ai_service):
        self.config = config
        self.ai_service = ai_service

    def build_report(self, metrics):
        timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        return "\n".join([
            "SIMULADOR DINÁMICO DE REDES · REPORTE DE EJECUCIÓN",
            "Universidad José Antonio Páez · Métodos Cuantitativos",
            f"Generado: {timestamp}",
            "",
            "[1] PARÁMETROS DE ENTRADA",
            f"    Tiempo de simulación .......... {metrics['sim_time']:.2f} s",
            f"    Tasa de llegada (lambda) ...... {metrics['arrival_rate']:.2f} paquetes/s",
            f"    Tasa de servicio (mu) ......... {metrics['service_rate']:.2f} paquetes/s",
            f"    Capacidad de buffer (S) ....... {metrics['buffer_capacity']} paquetes",
            f"    Umbral de reabastecimiento (s)  {metrics['replenish_threshold']} paquetes",
            f"    Lote de reposición (Q) ........ {metrics['batch_size']} paquetes",
            f"    Ponderación húngara (alpha) ... {metrics['alpha']:.2f}",
            "",
            "[2] MÉTRICAS DE DESEMPEÑO (SIMULADAS)",
            f"    Paquetes procesados ........... {metrics['packets_processed']}",
            f"    Paquetes perdidos ............. {metrics['packets_lost']}",
            f"    Tasa de pérdida ............... {metrics['loss_rate']:.2f} %",
            f"    Espera media en cola (Wq) ..... {metrics['wq']:.4f} s",
            f"    Tiempo medio en sistema (W) ... {metrics['w']:.4f} s",
            f"    Paquetes en sistema (L) ....... {metrics['l']:.2f}",
            f"    Paquetes en cola (Lq) ......... {metrics['lq']:.2f}",
            "",
            "[3] MODELO ANALÍTICO DE COLAS (M/M/1)",
            f"    Intensidad de tráfico (rho) ... {metrics['rho']:.3f}",
            f"    L teórico ..................... {self._format_finite(metrics['theoretical_l'])}",
            f"    Wq teórico .................... {self._format_finite(metrics['theoretical_wq'])} s",
            "",
            "[4] COSTOS DEL SISTEMA",
            f"    Costo de almacenamiento ....... ${metrics['holding_cost']:.2f}",
            f"    Costo de penalización ......... ${metrics['shortage_cost']:.2f}",
            f"    Costo global .................. ${metrics['total_cost']:.2f}",
        ])

    def _format_finite(self, value):
        if value == float("inf"):
            return "infinito (sistema saturado)"
        return f"{value:.2f}"

    def export(self, metrics):
        report = self.build_report(metrics)
        diagnosis = self.ai_service.analyze(report, metrics)
        content = f"{report}\n\n{'=' * 54}\n{diagnosis}\n"
        with open(self.config.report_txt_path, "w", encoding="utf-8") as handle:
            handle.write(content)
        self._write_document(metrics, diagnosis)
        print(f"[REPORTE] Guardado en {self.config.report_txt_path}")
        print()
        print(diagnosis)
        return content

    def _write_document(self, metrics, diagnosis):
        try:
            from docx import Document
            from docx.shared import Pt
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            return
        document = Document()
        title = document.add_heading("Informe Técnico · Simulador Dinámico de Redes", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle = document.add_paragraph("Universidad José Antonio Páez · Métodos Cuantitativos")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        document.add_heading("1. Parámetros y métricas obtenidas", level=1)
        table = document.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        header = table.rows[0].cells
        header[0].text = "Métrica"
        header[1].text = "Valor"
        for label, value in self._table_rows(metrics):
            row = table.add_row().cells
            row[0].text = label
            row[1].text = value

        document.add_heading("2. Diagnóstico automatizado", level=1)
        diagnosis_paragraph = document.add_paragraph(diagnosis)
        for run in diagnosis_paragraph.runs:
            run.font.size = Pt(9)

        document.save(self.config.report_docx_path)

    def _table_rows(self, metrics):
        return [
            ("Tiempo de simulación", f"{metrics['sim_time']:.2f} s"),
            ("Tasa de llegada (lambda)", f"{metrics['arrival_rate']:.2f} paq/s"),
            ("Tasa de servicio (mu)", f"{metrics['service_rate']:.2f} paq/s"),
            ("Capacidad de buffer (S)", f"{metrics['buffer_capacity']}"),
            ("Umbral de reabastecimiento (s)", f"{metrics['replenish_threshold']}"),
            ("Ponderación húngara (alpha)", f"{metrics['alpha']:.2f}"),
            ("Paquetes procesados", f"{metrics['packets_processed']}"),
            ("Paquetes perdidos", f"{metrics['packets_lost']}"),
            ("Tasa de pérdida", f"{metrics['loss_rate']:.2f} %"),
            ("Espera en cola (Wq)", f"{metrics['wq']:.4f} s"),
            ("Tiempo en sistema (W)", f"{metrics['w']:.4f} s"),
            ("Paquetes en sistema (L)", f"{metrics['l']:.2f}"),
            ("Intensidad de tráfico (rho)", f"{metrics['rho']:.3f}"),
            ("Costo de almacenamiento", f"${metrics['holding_cost']:.2f}"),
            ("Costo de penalización", f"${metrics['shortage_cost']:.2f}"),
            ("Costo global", f"${metrics['total_cost']:.2f}"),
        ]
